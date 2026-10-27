#!/usr/bin/env python3
"""PAKHAQS customer survey page: build data/pakhaqs.json.

Source: the "PAKHAQS Score Survey Summary report" workbook in the Store Assessment Drive folder (any
workbook whose name mentions PAKHAQS). Two sheets are read, recognised by their columns:

  outlet scores   one row per outlet x survey month: Phase, Month, Code, Outlet Code & Name, Store Type,
                  Total Sample, the seven PAKHAQS scores (Price Reasonable, Availability, Knowledge,
                  Housekeeping, Accurate, Quality, Speedy) and Avg Score (1-5)
  feedback        the latest month's "Category-wise Customer Feedback Summary": Rank, Feedback Category,
                  Customer Mentions, % of Valid Feedback, Key Interpretation (optional)

The newest workbook by Drive date is used; an unchanged file (same Drive id and modified time) is not
parsed again.

Output data/pakhaqs.json:
  dims     ["Price Reasonable", ...]        months  ["2024-06", ...]   phases {month: "Phase-10", ...}
  outlets  [[code, name, store type], ...]
  r        {o: [outlet index], m: [month index], n: [sample], s: [[7 scores]], a: [avg score]}
  feedback {title, lines: [...], rows: [[rank, category, mentions, share, interpretation], ...]}

Environment:
  PK_FOLDER_ID  Drive folder to search (default: the Store Assessment folder, SA_FOLDER_ID)
  PK_OUT        output file (default: data/pakhaqs.json)
Usage: python scripts/pakhaqs/refresh.py [--local FILE.xlsx]
"""
import datetime as dt
import json
import os
import re
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(HERE.parent / "network"))
from json_out import write_json  # noqa: E402

FOLDER = (os.environ.get("PK_FOLDER_ID") or os.environ.get("SA_FOLDER_ID") or "1TJ_c7VVyg6Qa_o0c62_LsZkd0sBHDEHF").strip()
OUT = Path(os.environ.get("PK_OUT") or ROOT / "data" / "pakhaqs.json")
NAME_RE = re.compile(r"pakhaqs", re.I)
DIMS = ["Price Reasonable", "Availability", "Knowledge", "Housekeeping", "Accurate", "Quality", "Speedy"]
MONTHS = {m: i for i, m in enumerate(("jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"), 1)}
CODE_RE = re.compile(r"^[A-Z]\d{3}$")


def log(msg):
    print(msg, flush=True)


def key(v):
    return re.sub(r"\s+", " ", str(v or "")).strip().lower()


def num(v):
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    try:
        return float(str(v).replace(",", "").strip())
    except (TypeError, ValueError):
        return None


def month(v):
    """"Sep'26", "Sept'26", a date -> "2026-09"."""
    if isinstance(v, (dt.date, dt.datetime)):
        return v.strftime("%Y-%m")
    m = re.match(r"\s*([A-Za-z]{3})[A-Za-z]*\W*(\d{2,4})", str(v or ""))
    if not m or m.group(1).lower() not in MONTHS:
        return None
    y = int(m.group(2))
    return f"{y + 2000 if y < 100 else y}-{MONTHS[m.group(1).lower()]:02d}"


def read_workbook(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        scores, feedback = None, None
        for ws in wb.worksheets:
            rows = ws.iter_rows(values_only=True)
            head = []
            for _ in range(12):
                try:
                    head.append(next(rows))
                except StopIteration:
                    break
            # outlet scores: the long sheet with Month, Code and Avg Score columns
            hi = next((i for i, r in enumerate(head) if {"month", "code", "avg score"} <= {key(v) for v in r if v is not None}), None)
            if hi is not None and scores is None:
                h = [key(v) for v in head[hi]]
                col = lambda name: h.index(name) if name in h else None  # noqa: E731
                ci = {"phase": col("phase"), "m": col("month"), "c": col("code"), "name": col("outlet code & name"), "type": col("store type"),
                      "n": col("total sample"), "a": col("avg score")}
                di = [col(key(d)) for d in DIMS]
                if None in (ci["m"], ci["c"], ci["a"]) or None in di:
                    continue
                out = []
                for r in list(head[hi + 1:]) + list(rows):
                    g = lambda i: r[i] if i is not None and i < len(r) else None  # noqa: E731
                    code, m = str(g(ci["c"]) or "").strip().upper(), month(g(ci["m"]))
                    a = num(g(ci["a"]))
                    if not CODE_RE.match(code) or not m or a is None:
                        continue
                    name = re.sub(rf"^{code}\s*[-–]?\s*", "", str(g(ci["name"]) or "").strip())
                    out.append((code, m, name, str(g(ci["type"]) or "").strip(), num(g(ci["n"])), [num(g(i)) for i in di], a, str(g(ci["phase"]) or "").strip()))
                scores = (ws.title, out)
                continue
            # feedback summary: Rank, Feedback Category, Customer Mentions ...
            fi = next((i for i, r in enumerate(head) if {"rank", "feedback category", "customer mentions"} <= {key(v) for v in r if v is not None}), None)
            if fi is not None and feedback is None:
                h = [key(v) for v in head[fi]]
                pick = lambda r, name: next((r[i] for i, x in enumerate(h) if x.startswith(name) and i < len(r)), None)  # noqa: E731
                lines = [str(r[0]).strip() for r in head[:fi] if r and r[0]]
                rows_out = []
                for r in list(head[fi + 1:]) + list(rows):
                    if num(pick(r, "rank")) is None or not pick(r, "feedback category"):
                        continue
                    rows_out.append([int(num(pick(r, "rank"))), str(pick(r, "feedback category")).strip(), num(pick(r, "customer mentions")), num(pick(r, "% of valid")), str(pick(r, "key interpretation") or "").strip()])
                feedback = {"sheet": ws.title, "title": lines[0] if lines else "", "lines": lines[1:], "rows": rows_out}
        return scores, feedback
    finally:
        wb.close()


def main():
    local = sys.argv[sys.argv.index("--local") + 1] if "--local" in sys.argv else None
    try:
        old = json.loads(OUT.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        old = {}
    if local:
        p = Path(local)
        items = [{"id": p.name, "name": p.name, "local": str(p), "modified": str(int(p.stat().st_mtime))}]
    else:
        import fetch_drive_data as drive
        log(f"Listing Drive folder {FOLDER}…")
        items = [i for i in drive.walk(FOLDER) if drive.is_spreadsheet(i) and NAME_RE.search(i["name"])]
        items.sort(key=drive.modified_sort_key, reverse=True)
    if not items:
        log("No PAKHAQS workbook in the Drive folder; pakhaqs.json left as is.")
        return 0
    item = items[0]
    fkey = f"{item['id']}|{item.get('modified', '')}"
    if (old.get("source") or {}).get("key") == fkey and old.get("schema") == 1:
        log(f"  keep  {item['name']} (unchanged)")
        return 0
    with tempfile.TemporaryDirectory(prefix="pk-") as work:
        path = item.get("local") or str(Path(work) / "pk.xlsx")
        if not item.get("local"):
            Path(path).write_bytes(drive.fetch_bytes(item))
        scores, feedback = read_workbook(path)
    if not scores or not scores[1]:
        log(f"::warning::{item['name']}: no outlet score sheet (Month, Code, Avg Score columns); pakhaqs.json left as is.")
        return 1
    sheet, rows = scores
    months = sorted({r[1] for r in rows})
    mi = {m: i for i, m in enumerate(months)}
    outlets, oi = [], {}
    for r in rows:
        if r[0] not in oi:
            oi[r[0]] = len(outlets)
            outlets.append([r[0], r[2], r[3]])
    phases = {}
    for r in rows:
        if r[7]:
            phases[r[1]] = r[7]
    rd = lambda v, n=3: None if v is None else round(v, n)  # noqa: E731
    payload = {
        "schema": 1, "generatedAt": dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds"),
        "source": {"key": fkey, "file": item["name"], "modified": item.get("modified", ""), "sheet": sheet, "rows": len(rows)},
        "dims": DIMS, "months": months, "phases": phases, "outlets": outlets,
        "r": {"o": [oi[r[0]] for r in rows], "m": [mi[r[1]] for r in rows], "n": [rd(r[4], 0) for r in rows],
              "s": [[rd(v) for v in r[5]] for r in rows], "a": [rd(r[6]) for r in rows]},
        "feedback": feedback,
    }
    if not write_json(OUT, payload):
        return 0
    log(f"Wrote {OUT} ({OUT.stat().st_size // 1024} KB): {len(rows)} outlet-months, {len(outlets)} outlets, {months[0]} to {months[-1]}"
        f"{', feedback ' + str(len(feedback['rows'])) + ' categories' if feedback else ''}.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
